"""Consolidates a UAE parent, its subsidiaries (including a sub-subsidiary) and its associates into group
financial statements under IFRS.

Steps: load each entity's trial balance, translate foreign subsidiaries into the
group currency, eliminate intercompany balances, sales, unrealised profit and the
investment in subsidiaries, recognise goodwill and non-controlling interests
(IFRS 10, IFRS 3, IAS 21), account for associates with the equity method (IAS 28),
and write an Excel workbook in which every step is a formula.

Run:  python consolidate.py
"""
import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

GROUP = "Illustrative Group"
GROUP_CCY = "USD"
UNIT = f"{GROUP_CCY} thousands"
TOLERANCE = 0.5  # largest intercompany mismatch accepted, in group currency

# Accounts the consolidation adds on top of the entity chart of accounts.
GROUP_ACCOUNTS = [(1700, "Goodwill", "A"), (3200, "Translation reserve", "E"), (3300, "Non-controlling interests", "E"),
                  (4200, "Share of profit of associates", "I")]
SC, RE, INV, GW, CTA, NCI = 3000, 3100, 1600, 1700, 3200, 3300
INV_ASSOC, DIV_ASSOC, SHARE_ASSOC = 1650, 4100, 4200
IC_REC, IC_PAY, IC_REV, IC_PUR, COS, INVENTORY = 1150, 2150, 4050, 5050, 5000, 1200


class ValidationError(Exception):
    pass


# ---------------------------------------------------------------- load and validate
def load_inputs(data_dir):
    d = Path(data_dir)
    tb = pd.read_csv(d / "trial_balances.csv")
    coa = pd.read_csv(d / "chart_of_accounts.csv")
    entities = pd.read_csv(d / "entities.csv")
    fx = pd.read_csv(d / "fx_rates.csv").set_index("currency")
    ic = pd.read_csv(d / "intercompany.csv")
    assoc = pd.read_csv(d / "associates.csv") if (d / "associates.csv").exists() else pd.DataFrame(
        columns=["entity", "name", "currency", "ownership", "investment_cost", "net_assets_at_acquisition",
                 "re_at_acquisition", "re_opening", "profit_for_year", "dividends_paid", "impairment"])

    tb["balance"] = tb.debit.fillna(0) - tb.credit.fillna(0)
    for ent, rows in tb.groupby("entity"):
        if round(rows.balance.sum(), 2) != 0:
            raise ValidationError(f"Trial balance for {ent} does not balance (difference {rows.balance.sum():,.2f})")
    unknown = sorted(set(tb.account_code) - set(coa.account_code))
    if unknown:
        raise ValidationError(f"Accounts not in the chart of accounts: {unknown}")
    missing = sorted(set(tb.entity) - set(entities.entity))
    if missing:
        raise ValidationError(f"Entities missing from entities.csv: {missing}")
    no_rate = sorted((set(entities.currency) | set(assoc.currency)) - set(fx.index))
    if no_rate:
        raise ValidationError(f"No exchange rates for: {no_rate}")
    if entities.iloc[0].ownership != 1:
        raise ValidationError("The first entity in entities.csv must be the parent (ownership 1)")
    if "held_by" not in entities:
        entities["held_by"] = [None] + [entities.entity.iloc[0]] * (len(entities) - 1)
    entities["held_by"] = entities.held_by.where(entities.held_by.notna(), None)
    entities.loc[0, "held_by"] = None
    ccy_of = entities.set_index("entity").currency
    effective = {entities.entity.iloc[0]: 1.0}
    for e in entities.iloc[1:].itertuples(index=False):
        if e.held_by not in effective:
            raise ValidationError(f"{e.entity} is held by {e.held_by}, which must be listed above it in entities.csv")
        effective[e.entity] = effective[e.held_by] * e.ownership
    entities["effective"] = entities.entity.map(effective)
    over = list(assoc.loc[assoc.ownership > 0.5, "entity"])
    if over:
        raise ValidationError(f"Associates owned over 50% are normally controlled (IFRS 10), so list them in "
                              f"entities.csv as subsidiaries: {over}")

    wide = tb.pivot_table(index="account_code", columns="entity", values="balance", aggfunc="sum").fillna(0)
    wide = wide.reindex(coa.account_code).fillna(0)[list(entities.entity)]
    return wide, coa, entities, fx, ic, assoc


# ---------------------------------------------------------------- consolidation (pandas; mirrors the Excel formulas)
def translate(wide, coa, entities, fx):
    """Translates each entity into the group currency: assets and liabilities at the closing rate,
    income and expenses at the average rate, share capital at the historical rate, opening retained
    earnings at the group-currency amount carried forward. The difference goes to the translation reserve."""
    types = coa.set_index("account_code").type
    out = {}
    for _, e in entities.iterrows():
        local = wide[e.entity]
        if e.currency == GROUP_CCY:
            usd = local.copy()
        else:
            r = fx.loc[e.currency]
            rate = types.map({"A": r.closing, "L": r.closing, "I": r.average, "X": r.average, "E": r.historical})
            usd = local * rate
            usd[RE] = -float(e.opening_re_group_currency)
        usd[CTA] = -usd.sum()
        out[e.entity] = usd
    return pd.DataFrame(out).fillna(0)


def equity_method(assoc, fx, ic):
    """IAS 28 equity method for each associate, in the group currency: cost, plus the group's share of
    post-acquisition reserves and of profit for the year, less dividends received, impairment and the
    group's share of unrealised profit on sales to or from the associate."""
    rows = []
    for _, a in assoc.iterrows():
        r = fx.loc[a.currency]
        own = a.ownership
        share_na = own * a.net_assets_at_acquisition * r.historical
        urp = ic[(ic.type == "associate_unrealised_profit") & ((ic.from_entity == a.entity) | (ic.to_entity == a.entity))]
        rows.append({
            "entity": a.entity, "name": a["name"], "ownership": own, "cost": a.investment_cost,
            "share_of_net_assets_at_acquisition": share_na, "goodwill": a.investment_cost - share_na,
            "post_acquisition_opening": own * (a.re_opening - a.re_at_acquisition) * r.historical,
            "share_of_profit": own * a.profit_for_year * r.average,
            "dividends_received": own * a.dividends_paid * r.average,
            "impairment": a.impairment,
            "urp_downstream": own * urp[urp.to_entity == a.entity].amount_group_currency.sum(),
            "urp_upstream": own * urp[urp.from_entity == a.entity].amount_group_currency.sum(),
        })
    em = pd.DataFrame(rows, columns=["entity", "name", "ownership", "cost", "share_of_net_assets_at_acquisition",
                                     "goodwill", "post_acquisition_opening", "share_of_profit", "dividends_received",
                                     "impairment", "urp_downstream", "urp_upstream"])
    em["carrying_amount"] = (em.cost + em.post_acquisition_opening + em.share_of_profit - em.dividends_received
                             - em.impairment - em.urp_downstream)
    return em


def eliminations(usd, entities, ic, em, fx):
    """Returns a list of (entry, description, account, amount) with debits positive.

    Goodwill and NCI use the group's effective interest, so a subsidiary held through another subsidiary gets
    the holder's NCI share of the cost of investment deducted (indirect holding). Goodwill of a foreign
    subsidiary is held in its own currency and retranslated at the closing rate (IAS 21.47).

    The investment cost is in the holding company's currency. A foreign holding company's investment sits in its
    translated trial balance at the closing rate, so that amount is removed; goodwill uses the cost at the
    historical rate, and the difference (the translation of the investment) comes out of the translation reserve
    for the group's share and out of NCI for the holder's NCI share."""
    j = []
    eff = entities.set_index("entity").effective
    ccy = entities.set_index("entity").currency
    for _, e in entities.iloc[1:].iterrows():
        r, rh = fx.loc[e.currency], fx.loc[ccy[e.held_by]]
        hist, clo = r.historical, r.closing
        own, eh, pre = e.effective, eff[e.held_by], e.pre_acquisition_re * hist
        inv_hist, inv = e.investment_cost * rh.historical, e.investment_cost * rh.closing
        sc = -usd.at[SC, e.entity]
        gw_hist = eh * inv_hist - own * (sc + pre)
        gw = gw_hist * clo / hist
        n = f"Investment in {e['name']}"
        j += [(n, SC, sc), (n, RE, pre), (n, GW, gw), (n, INV, -inv),
              (n, NCI, -((1 - own) * (sc + pre) - (1 - eh) * inv))]
        if e.currency != GROUP_CCY or ccy[e.held_by] != GROUP_CCY:
            j += [(n, CTA, gw_hist - gw + eh * (inv - inv_hist))]
        if own < 1:
            n = f"NCI share of {e['name']} post-acquisition reserves"
            post_re = -usd.at[RE, e.entity] - pre
            cta = usd.at[CTA, e.entity]
            j += [(n, RE, (1 - own) * post_re), (n, NCI, -(1 - own) * post_re),
                  (n, CTA, -(1 - own) * cta), (n, NCI, (1 - own) * cta)]
    for _, r in ic.iterrows():
        if r.type == "balance":
            n = f"Intercompany balance: {r.description}"
            j += [(n, IC_PAY, r.amount_group_currency), (n, IC_REC, -r.amount_group_currency)]
        elif r.type == "sale":
            n = f"Intercompany sale: {r.description}"
            j += [(n, IC_REV, r.amount_group_currency), (n, IC_PUR, -r.amount_group_currency)]
        elif r.type == "unrealised_profit":
            n = f"Unrealised profit in inventory: {r.description}"
            j += [(n, COS, r.amount_group_currency), (n, INVENTORY, -r.amount_group_currency)]
    for a in ([] if em is None else em.itertuples(index=False)):
        n = f"{a.name}: share of post-acquisition reserves at 1 January"
        j += [(n, INV_ASSOC, a.post_acquisition_opening), (n, RE, -a.post_acquisition_opening)]
        n = f"{a.name}: share of profit for the year"
        j += [(n, INV_ASSOC, a.share_of_profit), (n, SHARE_ASSOC, -a.share_of_profit)]
        n = f"{a.name}: dividends received reduce the investment"
        j += [(n, DIV_ASSOC, a.dividends_received), (n, INV_ASSOC, -a.dividends_received)]
        if a.impairment:
            n = f"{a.name}: impairment of the investment"
            j += [(n, SHARE_ASSOC, a.impairment), (n, INV_ASSOC, -a.impairment)]
        if a.urp_downstream:
            n = f"{a.name}: share of unrealised profit on sales to the associate"
            j += [(n, COS, a.urp_downstream), (n, INV_ASSOC, -a.urp_downstream)]
        if a.urp_upstream:
            n = f"{a.name}: share of unrealised profit on goods bought from the associate"
            j += [(n, SHARE_ASSOC, a.urp_upstream), (n, INVENTORY, -a.urp_upstream)]
    return j


def consolidate(wide, coa, entities, fx, ic, assoc=None):
    usd = translate(wide, coa, entities, fx)
    usd = usd.reindex(list(usd.index) + [a[0] for a in GROUP_ACCOUNTS if a[0] not in usd.index]).fillna(0)
    em = equity_method(assoc, fx, ic) if assoc is not None and len(assoc) else None
    journal = eliminations(usd, entities, ic, em, fx)
    elim = pd.Series(0.0, index=usd.index)
    for _, acct, amt in journal:
        elim[acct] += amt
    for entry in {j[0] for j in journal}:
        if abs(sum(a for n, _, a in journal if n == entry)) > 1e-6:
            raise ValidationError(f"Elimination entry does not balance: {entry}")
    consolidated = usd.sum(axis=1) + elim
    for acct, label in ((IC_REC, "receivables"), (IC_PAY, "payables"), (IC_REV, "revenue"), (IC_PUR, "purchases")):
        if abs(consolidated[acct]) > TOLERANCE:
            raise ValidationError(f"Intercompany {label} do not eliminate: {consolidated[acct]:,.2f} left over. "
                                  f"Check intercompany.csv against the entity trial balances.")
    if abs(consolidated.get(DIV_ASSOC, 0)) > TOLERANCE:
        raise ValidationError(f"Dividend income from associates does not agree to associates.csv: "
                              f"{-consolidated[DIV_ASSOC]:,.2f} left over.")
    return usd, journal, elim, consolidated, em


def summarise(usd, entities, consolidated):
    c = -consolidated
    profit = c[[a for a in c.index if 4000 <= a < 9000]].sum()
    nci_profit = sum((1 - e.effective) * -usd[e.entity][[a for a in usd.index if 4000 <= a < 9000]].sum()
                     for _, e in entities.iloc[1:].iterrows())
    s = {"Revenue": c[4000] + c[4050], "Profit for the year": profit, "Attributable to owners": profit - nci_profit,
         "Attributable to NCI": nci_profit, "Share of profit of associates": c.get(SHARE_ASSOC, 0),
         "Goodwill": -c[GW], "Investment in associates": -c.get(INV_ASSOC, 0),
         "Total assets": -c[[a for a in c.index if a < 2000]].sum(),
         "Equity attributable to owners": c[SC] + c[RE] + c[CTA] + profit - nci_profit,
         "Non-controlling interests": c[NCI] + nci_profit,
         "Total liabilities": c[[a for a in c.index if 2000 <= a < 3000]].sum()}
    s["Total equity and liabilities"] = s["Equity attributable to owners"] + s["Non-controlling interests"] + s["Total liabilities"]
    return {k: round(float(v), 2) for k, v in s.items()}


# ---------------------------------------------------------------- Excel output
FONT = "Arial"
HFILL = PatternFill("solid", fgColor="1F4E78")
TOP = Border(top=Side(style="thin"))
DBL = Border(top=Side(style="thin"), bottom=Side(style="double"))
NUM = '#,##0;(#,##0);"-"'
RATE = "0.0000"


def _f(bold=False, color="000000", italic=False, size=10):
    return Font(name=FONT, bold=bold, color=color, italic=italic, size=size)


def _sheet(wb, title, heading, sub=UNIT):
    ws = wb.create_sheet(title)
    ws["B1"] = GROUP; ws["B1"].font = _f(True, size=14)
    ws["B2"] = heading; ws["B2"].font = _f(True, size=12)
    ws["B3"] = sub; ws["B3"].font = _f(italic=True, color="595959")
    ws.column_dimensions["A"].width = 2
    ws.sheet_view.showGridLines = False
    return ws


def _header(ws, row, labels, first_col=2, widths=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=first_col + i, value=lab)
        c.font = _f(True, "FFFFFF"); c.fill = HFILL
        c.alignment = Alignment(horizontal="left" if i < 2 else "right", wrap_text=True)
        if widths: ws.column_dimensions[c.column_letter].width = widths[i]


def write_workbook(wide, coa, entities, fx, ic, journal, out_path, assoc=None):
    wb = Workbook(); wb.remove(wb.active)
    ents = list(entities.entity)
    accounts = list(coa.itertuples(index=False)) + [type(coa.iloc[0:1].itertuples(index=False).__next__())(*a)
                                                     for a in GROUP_ACCOUNTS]

    # Inputs: entities, FX, intercompany
    inp = _sheet(wb, "Inputs", "Group structure, exchange rates and intercompany data")
    _header(inp, 5, ["Entity", "Name", "Currency", "Ownership", "Investment cost (holder's currency)",
                     "Pre-acq. retained earnings (local)", "Opening RE (USD)", "Held by", "Group effective interest"],
            widths=[10, 28, 10, 11, 12, 14, 12, 10, 12])
    ent_row = {}
    for i, e in enumerate(entities.itertuples(index=False)):
        r = 6 + i; ent_row[e.entity] = r
        for j, v in enumerate([e.entity, e.name, e.currency, e.ownership, e.investment_cost, e.pre_acquisition_re,
                               None if pd.isna(e.opening_re_group_currency) else e.opening_re_group_currency,
                               e.held_by]):
            c = inp.cell(row=r, column=2 + j, value=v); c.font = _f(color="0000FF" if 3 <= j < 7 else "000000")
            c.number_format = "0%" if j == 3 else NUM
            if j == 7: c.alignment = Alignment(horizontal="right")
        c = inp.cell(row=r, column=10, value=1 if i == 0 else f"=E{r}*J{ent_row[e.held_by]}")
        c.number_format = "0%"; c.font = _f()
    fx_top = 6 + len(ents) + 2
    _header(inp, fx_top, ["Currency", "", "Closing", "Average", "Historical"])
    fx_row = {}
    for i, (ccy, r_) in enumerate(fx.iterrows()):
        r = fx_top + 1 + i; fx_row[ccy] = r
        inp.cell(row=r, column=2, value=ccy).font = _f()
        for j, k in enumerate(("closing", "average", "historical")):
            c = inp.cell(row=r, column=4 + j, value=float(r_[k])); c.font = _f(color="0000FF"); c.number_format = RATE
    ic_top = fx_top + len(fx) + 3
    _header(inp, ic_top, ["Type", "Description", "From", "To", "Amount (USD)"])
    inp.column_dimensions["C"].width = 46
    ic_row = []
    for i, x in enumerate(ic.itertuples(index=False)):
        r = ic_top + 1 + i; ic_row.append(r)
        for j, v in enumerate([x.type, x.description, x.from_entity, x.to_entity, x.amount_group_currency]):
            c = inp.cell(row=r, column=2 + j, value=v); c.font = _f(color="0000FF" if j == 4 else "000000")
            if j == 4: c.number_format = NUM

    # Trial balances in local currency
    tbs = _sheet(wb, "Trial Balances", "Entity trial balances in local currency (debit +, credit -)", "thousands of local currency")
    _header(tbs, 5, ["Code", "Account", "Type"] + [f"{e} ({entities.set_index('entity').currency[e]})" for e in ents],
            widths=[8, 30, 6] + [13] * len(ents))
    tb_row = {}
    for i, a in enumerate(coa.itertuples(index=False)):
        r = 6 + i; tb_row[a.account_code] = r
        for j, v in enumerate([a.account_code, a.account_name, a.type]):
            tbs.cell(row=r, column=2 + j, value=v).font = _f()
        for k, e in enumerate(ents):
            c = tbs.cell(row=r, column=5 + k, value=float(wide.at[a.account_code, e]))
            c.font = _f(color="0000FF"); c.number_format = NUM
    last_tb = 6 + len(coa) - 1
    tbs.cell(row=last_tb + 1, column=3, value="Check (should be zero)").font = _f(True)
    for k in range(len(ents)):
        col = tbs.cell(row=6, column=5 + k).column_letter
        c = tbs[f"{col}{last_tb + 1}"]; c.value = f"=SUM({col}6:{col}{last_tb})"; c.number_format = NUM
        c.font = _f(True); c.border = TOP

    # Translation into USD
    tr = _sheet(wb, "Translation", "Translation into USD")
    tr_cols = {}
    widths = [8, 30]
    labels = ["Code", "Account"]
    for e in ents:
        labels += [f"{e} local", f"{e} rate", f"{e} USD"]; widths += [11, 9, 11]
    _header(tr, 5, labels, widths=widths)
    tr_row = {}
    rows = [(a.account_code, a.account_name, a.type) for a in coa.itertuples(index=False)] + [(CTA, "Translation reserve", "E")]
    for i, (code, name, typ) in enumerate(rows):
        r = 6 + i; tr_row[code] = r
        tr.cell(row=r, column=2, value=code).font = _f(); tr.cell(row=r, column=3, value=name).font = _f()
    for k, e in enumerate(ents):
        lc, rc, uc = (tr.cell(row=5, column=4 + 3 * k + m).column_letter for m in range(3))
        tr_cols[e] = uc
        tbcol = tbs.cell(row=5, column=5 + k).column_letter
        fr = fx_row[entities.set_index("entity").currency[e]]
        for code, name, typ in rows[:-1]:
            r = tr_row[code]
            tr[f"{lc}{r}"] = f"='Trial Balances'!{tbcol}{tb_row[code]}"
            basis = {"A": "D", "L": "D", "I": "E", "X": "E", "E": "F"}[typ]
            tr[f"{rc}{r}"] = f"=Inputs!${basis}${fr}"
            if code == RE and entities.set_index("entity").currency[e] != GROUP_CCY:
                tr[f"{uc}{r}"] = f"=-Inputs!$H${ent_row[e]}"
                tr[f"{rc}{r}"] = f"=IF({lc}{r}=0,0,{uc}{r}/{lc}{r})"
            else:
                tr[f"{uc}{r}"] = f"={lc}{r}*{rc}{r}"
        r = tr_row[CTA]
        tr[f"{uc}{r}"] = f"=-SUM({uc}6:{uc}{r - 1})"
        for code, *_ in rows:
            for col, fmt, color in ((lc, NUM, "008000"), (rc, RATE, "008000"), (uc, NUM, "000000")):
                cell = tr[f"{col}{tr_row[code]}"]
                if cell.value is not None: cell.number_format = fmt; cell.font = _f(color=color)
        tr[f"{uc}{r}"].font = _f(True); tr[f"{uc}{r}"].border = TOP
    tr.cell(row=tr_row[CTA] + 2, column=2, value="Assets and liabilities at closing rate, income and expenses at average "
            "rate, share capital at historical rate, opening retained earnings at the USD amount carried forward. "
            "The translation reserve balances the translated trial balance.").font = _f(italic=True, color="595959")

    # Associates: IAS 28 equity method schedule
    assoc = assoc if assoc is not None else pd.DataFrame()
    acol, ar = {}, {}
    if len(assoc):
        asx = _sheet(wb, "Associates", "Associates: equity method (IAS 28)")
        _header(asx, 5, ["", ""] + list(assoc.entity) + ["Total"], widths=[66, 2] + [20] * (len(assoc) + 1))
        tot = asx.cell(row=5, column=4 + len(assoc)).column_letter
        for k, e in enumerate(assoc.entity):
            acol[e] = asx.cell(row=5, column=4 + k).column_letter
        lines = [  # key, label, kind ('in' input, 'f' formula, 'h' heading), format, in total column
            ("name", "Name", "in", None, False), ("ccy", "Currency", "in", None, False),
            ("own", "Ownership", "in", "0%", False), ("cost", "Cost of investment (USD)", "in", NUM, True),
            ("na", "Net assets at acquisition (local currency)", "in", NUM, False),
            ("rea", "Retained earnings at acquisition (local currency)", "in", NUM, False),
            ("reo", "Retained earnings at 1 January 2025 (local currency)", "in", NUM, False),
            ("pft", "Profit for the year (local currency)", "in", NUM, False),
            ("div", "Dividends paid by the associate (local currency)", "in", NUM, False),
            ("imp", "Impairment loss from IAS 36 test (USD)", "in", NUM, True),
            ("hist", "Historical rate", "f", RATE, False), ("avg", "Average rate", "f", RATE, False),
            ("clo", "Closing rate", "f", RATE, False),
            (None, None, None, None, None),
            ("h1", "Equity method (USD)", "h", None, False),
            ("sna", "Group share of net assets at acquisition", "f", NUM, True),
            ("gw", "Goodwill included in the investment (cost less share of net assets)", "f", NUM, True),
            (None, None, None, None, None),
            ("c0", "Cost of investment", "f", NUM, True),
            ("post", "Share of post-acquisition retained earnings to 1 January", "f", NUM, True),
            ("open", "Carrying amount at 1 January 2025", "f", NUM, True),
            ("sop", "Share of profit for the year", "f", NUM, True),
            ("dvr", "Dividends received (reduce the investment)", "f", NUM, True),
            ("impn", "Impairment loss", "f", NUM, True),
            ("urpd", "Group share of unrealised profit on sales to the associate (downstream)", "f", NUM, True),
            ("close", "Carrying amount at 31 December 2025", "f", NUM, True),
            (None, None, None, None, None),
            ("h2", "Reconciliation to the associate's net assets (IFRS 12)", "h", None, False),
            ("nac", "Associate's net assets at 31 December (local currency)", "f", NUM, False),
            ("snac", "Group share at closing rate", "f", NUM, True),
            ("gw2", "Goodwill", "f", NUM, True), ("imp2", "Impairment loss", "f", NUM, True),
            ("urp2", "Unrealised profit (downstream)", "f", NUM, True),
            ("fxd", "Exchange difference (zero for USD and the pegged AED)", "f", NUM, True),
            ("close2", "Carrying amount at 31 December 2025", "f", NUM, True),
            (None, None, None, None, None),
            ("urpu", "Memo: share of upstream unrealised profit (reduces group inventory)", "f", NUM, True),
        ]
        for i, (key, *_r) in enumerate(lines):
            if key: ar[key] = 6 + i
        ic_last = ic_row[-1] if ic_row else ic_top + 1
        for i, (key, label, kind, fmt, in_total) in enumerate(lines):
            r = 6 + i
            if key is None: continue
            asx[f"B{r}"] = label
            asx[f"B{r}"].font = _f(kind == "h" or key in ("open", "close", "close2"))
            if kind == "h": continue
            for a in assoc.itertuples(index=False):
                col = acol[a.entity]; R = lambda k: f"{col}{ar[k]}"
                fr = fx_row[a.currency]
                src = {"name": a.name, "ccy": a.currency, "own": a.ownership, "cost": a.investment_cost,
                       "na": a.net_assets_at_acquisition, "rea": a.re_at_acquisition, "reo": a.re_opening,
                       "pft": a.profit_for_year, "div": a.dividends_paid, "imp": a.impairment}
                ic_sum = (f"SUMIFS(Inputs!$F${ic_top + 1}:$F${ic_last},Inputs!$B${ic_top + 1}:$B${ic_last},"
                          f"\"associate_unrealised_profit\",Inputs!${{}}${ic_top + 1}:${{}}${ic_last},\"{a.entity}\")")
                f = {"hist": f"=Inputs!$F${fr}", "avg": f"=Inputs!$E${fr}", "clo": f"=Inputs!$D${fr}",
                     "sna": f"={R('own')}*{R('na')}*{R('hist')}", "gw": f"={R('cost')}-{R('sna')}",
                     "c0": f"={R('cost')}", "post": f"={R('own')}*({R('reo')}-{R('rea')})*{R('hist')}",
                     "open": f"={R('c0')}+{R('post')}", "sop": f"={R('own')}*{R('pft')}*{R('avg')}",
                     "dvr": f"=-{R('own')}*{R('div')}*{R('avg')}", "impn": f"=-{R('imp')}",
                     "urpd": f"=-{R('own')}*" + ic_sum.format("E", "E"),
                     "close": f"=SUM({R('open')}:{R('urpd')})",
                     "nac": f"={R('na')}+{R('reo')}-{R('rea')}+{R('pft')}-{R('div')}",
                     "snac": f"={R('own')}*{R('nac')}*{R('clo')}", "gw2": f"={R('gw')}", "imp2": f"={R('impn')}",
                     "urp2": f"={R('urpd')}", "fxd": f"={R('close')}-SUM({R('snac')}:{R('urp2')})",
                     "close2": f"=SUM({R('snac')}:{R('fxd')})",
                     "urpu": f"={R('own')}*" + ic_sum.format("D", "D")}
                c = asx[R(key)]
                if kind == "in":
                    c.value = src[key]; c.font = _f(color="0000FF")
                else:
                    c.value = f[key]; c.font = _f(key in ("open", "close", "close2"),
                                                 color="008000" if key in ("hist", "avg", "clo") else "000000")
                if fmt: c.number_format = fmt
                c.alignment = Alignment(horizontal="right")
                if key in ("open", "close", "close2", "c0", "snac"): pass
                if key in ("open", "close", "close2"): c.border = TOP if key == "open" else DBL
            if in_total:
                c = asx[f"{tot}{r}"]; c.value = f"=SUM({acol[assoc.entity.iloc[0]]}{r}:{acol[assoc.entity.iloc[-1]]}{r})"
                c.number_format = NUM; c.font = _f(True)
        asx.cell(row=6 + len(lines) + 1, column=2, value="The group does not combine an associate line by line. "
                 "It carries one investment line, measured at cost plus its share of post-acquisition profits, less "
                 "dividends, impairment and its share of unrealised profit. Balances with associates are not "
                 "eliminated, because associates are outside the group.").font = _f(italic=True, color="595959")
        asx.cell(row=6 + len(lines) + 1, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        asx.row_dimensions[6 + len(lines) + 1].height = 54

    # Eliminations journal
    el = _sheet(wb, "Eliminations", "Consolidation adjustments (debit +, credit -)")
    _header(el, 5, ["#", "Description", "Code", "Account", "Amount"], widths=[5, 72, 8, 30, 12])
    name_of = {a[0]: a[1] for a in [(x.account_code, x.account_name) for x in coa.itertuples(index=False)] + [(g[0], g[1]) for g in GROUP_ACCOUNTS]}
    r = 6; entry_no = {}
    ent_by = entities.set_index("entity")
    sub_iter = list(entities.iloc[1:].itertuples(index=False))

    def usd(code, ent):
        return f"Translation!{tr_cols[ent]}{tr_row[code]}"
    formulas = []
    for e in sub_iter:
        er = ent_row[e.entity]; fr = fx_row[e.currency]; hr = fx_row[ent_by.currency[e.held_by]]
        own, eh = f"Inputs!$J${er}", f"Inputs!$J${ent_row[e.held_by]}"
        inv, inv_hist = f"(Inputs!$F${er}*Inputs!$D${hr})", f"(Inputs!$F${er}*Inputs!$F${hr})"
        pre = f"(Inputs!$G${er}*Inputs!$F${fr})"
        n = f"Investment in {e.name}"
        sc = f"(-{usd(SC, e.entity)})"
        gw_hist = f"({eh}*{inv_hist}-{own}*({sc}+{pre}))"
        formulas += [(n, SC, f"={sc}"), (n, RE, f"={pre}"), (n, GW, f"={gw_hist}*Inputs!$D${fr}/Inputs!$F${fr}"),
                     (n, INV, f"=-{inv}"), (n, NCI, f"=-((1-{own})*({sc}+{pre})-(1-{eh})*{inv})")]
        if e.currency != GROUP_CCY or ent_by.currency[e.held_by] != GROUP_CCY:
            formulas += [(n, CTA, f"={gw_hist}*(1-Inputs!$D${fr}/Inputs!$F${fr})+{eh}*({inv}-{inv_hist})")]
        if e.effective < 1:
            n = f"NCI share of {e.name} post-acquisition reserves"
            post = f"(-{usd(RE, e.entity)}-{pre})"; cta = usd(CTA, e.entity)
            formulas += [(n, RE, f"=(1-{own})*{post}"), (n, NCI, f"=-(1-{own})*{post}"),
                         (n, CTA, f"=-(1-{own})*{cta}"), (n, NCI, f"=(1-{own})*{cta}")]
    for x, ir in zip(ic.itertuples(index=False), ic_row):
        amt = f"Inputs!$F${ir}"
        if x.type == "balance":
            n = f"Intercompany balance: {x.description}"; formulas += [(n, IC_PAY, f"={amt}"), (n, IC_REC, f"=-{amt}")]
        elif x.type == "sale":
            n = f"Intercompany sale: {x.description}"; formulas += [(n, IC_REV, f"={amt}"), (n, IC_PUR, f"=-{amt}")]
        elif x.type == "unrealised_profit":
            n = f"Unrealised profit in inventory: {x.description}"; formulas += [(n, COS, f"={amt}"), (n, INVENTORY, f"=-{amt}")]
    for a in assoc.itertuples(index=False) if len(assoc) else []:
        A = lambda k: f"Associates!{acol[a.entity]}{ar[k]}"
        n = f"{a.name}: share of post-acquisition reserves at 1 January"
        formulas += [(n, INV_ASSOC, f"={A('post')}"), (n, RE, f"=-{A('post')}")]
        n = f"{a.name}: share of profit for the year"
        formulas += [(n, INV_ASSOC, f"={A('sop')}"), (n, SHARE_ASSOC, f"=-{A('sop')}")]
        n = f"{a.name}: dividends received reduce the investment"
        formulas += [(n, DIV_ASSOC, f"=-{A('dvr')}"), (n, INV_ASSOC, f"={A('dvr')}")]
        if a.impairment:
            n = f"{a.name}: impairment of the investment"
            formulas += [(n, SHARE_ASSOC, f"=-{A('impn')}"), (n, INV_ASSOC, f"={A('impn')}")]
        if ic[(ic.type == "associate_unrealised_profit") & (ic.to_entity == a.entity)].shape[0]:
            n = f"{a.name}: share of unrealised profit on sales to the associate"
            formulas += [(n, COS, f"=-{A('urpd')}"), (n, INV_ASSOC, f"={A('urpd')}")]
        if ic[(ic.type == "associate_unrealised_profit") & (ic.from_entity == a.entity)].shape[0]:
            n = f"{a.name}: share of unrealised profit on goods bought from the associate"
            formulas += [(n, SHARE_ASSOC, f"={A('urpu')}"), (n, INVENTORY, f"=-{A('urpu')}")]
    for n, code, f in formulas:
        entry_no.setdefault(n, len(entry_no) + 1)
        for j, v in enumerate([entry_no[n], n, code, name_of[code], f]):
            c = el.cell(row=r, column=2 + j, value=v); c.font = _f(color="000000" if j < 4 else "008000")
            if j == 4: c.number_format = NUM
        r += 1
    el_last = r - 1
    el.cell(row=r, column=5, value="Total (should be zero)").font = _f(True)
    c = el.cell(row=r, column=6, value=f"=ROUND(SUM(F6:F{el_last}),2)"); c.number_format = NUM; c.font = _f(True); c.border = TOP

    # Consolidation worksheet
    cs = _sheet(wb, "Consolidation", "Consolidation worksheet (debit +, credit -)")
    _header(cs, 5, ["Code", "Account"] + [f"{e} (USD)" for e in ents] + ["Aggregate", "Eliminations", "Consolidated"],
            widths=[8, 30] + [13] * (len(ents) + 3))
    cs_row = {}
    all_rows = [(a.account_code, a.account_name) for a in coa.itertuples(index=False)] + [(g[0], g[1]) for g in GROUP_ACCOUNTS]
    all_rows.sort()
    ecol = [cs.cell(row=5, column=4 + k).column_letter for k in range(len(ents))]
    agg, elc, con = (cs.cell(row=5, column=4 + len(ents) + m).column_letter for m in range(3))
    for i, (code, name) in enumerate(all_rows):
        r = 6 + i; cs_row[code] = r
        cs.cell(row=r, column=2, value=code).font = _f(); cs.cell(row=r, column=3, value=name).font = _f()
        for k, e in enumerate(ents):
            cs[f"{ecol[k]}{r}"] = f"=Translation!{tr_cols[e]}{tr_row[code]}" if code in tr_row else 0
        cs[f"{agg}{r}"] = f"=SUM({ecol[0]}{r}:{ecol[-1]}{r})"
        cs[f"{elc}{r}"] = f"=SUMIFS(Eliminations!$F$6:$F${el_last},Eliminations!$D$6:$D${el_last},$B{r})"
        cs[f"{con}{r}"] = f"={agg}{r}+{elc}{r}"
        for col in ecol + [agg, elc, con]:
            c = cs[f"{col}{r}"]; c.number_format = NUM
            c.font = _f(color="008000" if col in ecol or col == elc else "000000", bold=col == con)
    cs_last = 6 + len(all_rows) - 1
    cs.cell(row=cs_last + 1, column=3, value="Check (should be zero)").font = _f(True)
    for col in ecol + [agg, elc, con]:
        c = cs[f"{col}{cs_last + 1}"]; c.value = f"=ROUND(SUM({col}6:{col}{cs_last}),2)"; c.number_format = NUM
        c.font = _f(True); c.border = TOP

    # Presentation statements
    C = lambda code: f"Consolidation!{con}{cs_row[code]}"
    def T(code, ent): return f"Translation!{tr_cols[ent]}{tr_row[code]}"
    is_ = _sheet(wb, "Consolidated IS", "Consolidated income statement for the year ended 31 December 2025")
    is_.column_dimensions["B"].width = 44; is_.column_dimensions["C"].width = 14
    nci_terms = "+".join(f"(1-Inputs!$J${ent_row[e.entity]})*-SUM(Translation!{tr_cols[e.entity]}{tr_row[4000]}:"
                         f"{tr_cols[e.entity]}{tr_row[8000]})" for e in sub_iter)
    is_lines = [
        ("Revenue", f"=-({C(4000)}+{C(4050)})", None),
        ("Cost of sales", f"=-({C(5000)}+{C(5050)})", None),
        ("Gross profit", "=C6+C7", TOP),
        ("Operating expenses", f"=-{C(6000)}", None),
        ("Operating profit", "=C8+C9", TOP),
        ("Share of profit of associates (equity method)", f"=-{C(SHARE_ASSOC)}", None),
        ("Dividend income from associates (eliminated)", f"=-{C(DIV_ASSOC)}", None),
        ("Interest expense", f"=-{C(7000)}", None),
        ("Profit before tax", "=SUM(C10:C13)", TOP),
        ("Income tax expense", f"=-{C(8000)}", None),
        ("Profit for the year", "=C14+C15", DBL),
        (None, None, None),
        ("Attributable to owners of the parent", "=C16-C19", None),
        ("Attributable to non-controlling interests", f"={nci_terms}", None),
    ]
    _header(is_, 5, ["", "2025"], widths=None)
    for i, (lab, f, border) in enumerate(is_lines):
        r = 6 + i
        if lab is None: continue
        is_[f"B{r}"] = lab; is_[f"C{r}"] = f; is_[f"C{r}"].number_format = NUM
        bold = border is not None
        is_[f"B{r}"].font = _f(bold); is_[f"C{r}"].font = _f(bold)
        if border: is_[f"C{r}"].border = border
    is_["B21"] = ("Unrealised profit is eliminated in full against the seller (Sub One, 100% owned), so it has no NCI effect.")
    is_["B22"] = ("Share of profit of associates is after tax and after the impairment loss (see the Associates sheet). "
                  "The group's share of unrealised profit on sales to associates is charged to cost of sales.")
    for c in ("B21", "B22"): is_[c].font = _f(italic=True, color="595959")

    bs = _sheet(wb, "Consolidated BS", "Consolidated balance sheet at 31 December 2025")
    bs.column_dimensions["B"].width = 44; bs.column_dimensions["C"].width = 14
    _header(bs, 5, ["", "2025"])
    bs_lines = [
        ("Goodwill", f"={C(GW)}", None), ("Property, plant and equipment", f"={C(1500)}", None),
        ("Investments in associates (equity method)", f"={C(INV_ASSOC)}", None),
        ("Total non-current assets", "=SUM(C6:C8)", TOP),
        ("Inventory", f"={C(1200)}", None), ("Trade receivables", f"={C(1100)}", None),
        ("Intercompany receivable (eliminated)", f"={C(IC_REC)}", None), ("Cash", f"={C(1000)}", None),
        ("Total current assets", "=SUM(C10:C13)", TOP),
        ("Total assets", "=C9+C14", DBL),
        (None, None, None),
        ("Share capital", f"=-{C(SC)}", None),
        ("Retained earnings", f"=-{C(RE)}+'Consolidated IS'!C18", None),
        ("Translation reserve", f"=-{C(CTA)}", None),
        ("Equity attributable to owners of the parent", "=SUM(C17:C19)", TOP),
        ("Non-controlling interests", f"=-{C(NCI)}+'Consolidated IS'!C19", None),
        ("Total equity", "=C20+C21", TOP),
        ("Borrowings", f"=-{C(2500)}", None), ("Trade payables", f"=-{C(2000)}", None),
        ("Intercompany payable (eliminated)", f"=-{C(IC_PAY)}", None),
        ("Total liabilities", "=SUM(C23:C25)", TOP),
        ("Total equity and liabilities", "=C22+C26", DBL),
        (None, None, None),
        ("Check: assets less equity and liabilities", "=ROUND(C15-C27,2)", None),
    ]
    for i, (lab, f, border) in enumerate(bs_lines):
        r = 6 + i
        if lab is None: continue
        bs[f"B{r}"] = lab; bs[f"C{r}"] = f; bs[f"C{r}"].number_format = NUM
        bold = border is not None
        bs[f"B{r}"].font = _f(bold); bs[f"C{r}"].font = _f(bold)
        if border: bs[f"C{r}"].border = border

    # Checks
    ck = _sheet(wb, "Checks", "Checks (all should be zero)")
    ck.column_dimensions["B"].width = 52; ck.column_dimensions["C"].width = 14
    checks = [("Entity trial balances balance",
               "+".join(f"ABS('Trial Balances'!{tbs.cell(row=5, column=5 + k).column_letter}{last_tb + 1})" for k in range(len(ents)))),
              ("Elimination journal balances", f"ABS(Eliminations!F{el_last + 1})"),
              ("Consolidation worksheet balances", f"ABS(Consolidation!{con}{cs_last + 1})"),
              ("Intercompany receivables and payables fully eliminated", f"ABS({C(IC_REC)})+ABS({C(IC_PAY)})"),
              ("Intercompany sales and purchases fully eliminated", f"ABS({C(IC_REV)})+ABS({C(IC_PUR)})"),
              ("Investment in subsidiaries fully eliminated", f"ABS({C(INV)})"),
              ("Balance sheet balances", "ABS('Consolidated BS'!C29)")]
    if len(assoc):
        tot = asx.cell(row=5, column=4 + len(assoc)).column_letter
        checks[-1:-1] = [("Dividends from associates fully eliminated", f"ABS({C(DIV_ASSOC)})"),
                         ("Investments in associates agree to the equity method schedule",
                          f"ABS({C(INV_ASSOC)}-Associates!{tot}{ar['close']})"),
                         ("Associates reconcile to share of net assets", f"ABS(Associates!{tot}{ar['fxd']})")]
    for i, (lab, f) in enumerate(checks):
        ck[f"B{6 + i}"] = lab; ck[f"B{6 + i}"].font = _f()
        ck[f"C{6 + i}"] = f"=ROUND({f},2)"; ck[f"C{6 + i}"].number_format = NUM; ck[f"C{6 + i}"].font = _f()
    r = 6 + len(checks) + 1
    ck[f"B{r}"] = "All checks pass?"; ck[f"B{r}"].font = _f(True)
    ck[f"C{r}"] = f'=IF(SUM(C6:C{r - 2})=0,"OK","ERROR")'; ck[f"C{r}"].font = _f(True)

    _ifrs_sheet(wb, entities, assoc, ent_row, acol, ar)
    order = ["Consolidated IS", "Consolidated BS", "IFRS Treatment", "Consolidation", "Eliminations"] + \
        (["Associates"] if len(assoc) else []) + ["Translation", "Trial Balances", "Inputs", "Checks"]
    wb._sheets = [wb[n] for n in order]
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    wb.save(out_path)


# ---------------------------------------------------------------- IFRS treatment sheet
IFRS_STEPS = [
    ("Control (subsidiaries)",
     "An investor controls an investee when it has power over it, is exposed to variable returns from it and can use "
     "its power to affect those returns. Subsidiaries One and Two are controlled directly through a majority of "
     "the voting rights, Subsidiary Three through Subsidiary One and Subsidiary Four through Subsidiary Two. "
     "Each is consolidated from the date "
     "control was obtained.", "IFRS 10.6-7, B35"),
    ("Significant influence (associates)",
     "Significant influence is the power to take part in financial and operating policy decisions without control. "
     "A holding of 20% or more of the voting power is presumed to give it, and board representation is typical "
     "evidence. Both associates are held at 25% to 30%.", "IAS 28.3, 28.5-6"),
    ("Uniform policies and reporting date",
     "All entities use the group's IFRS accounting policies and a 31 December year end, so no alignment "
     "adjustments are needed. Entities in countries that do not use IFRS locally (such as the US) report to the "
     "group on IFRS policies. A different year end would need adjusting for significant transactions and may "
     "differ by no more than three months.", "IFRS 10.B87, B92-B93; IAS 28.33-36"),
    ("Foreign currency translation",
     "Subsidiary Two (EUR) is translated into USD: assets and liabilities at the closing rate, income and "
     "expenses at the average rate (an approximation of transaction-date rates) and share capital at the "
     "historical rate. The difference goes to the translation reserve in other comprehensive income, and the "
     "NCI's share is allocated to NCI. Subsidiary Three (GBP) is translated the same way, and its goodwill is "
     "treated as its own asset: held in GBP and retranslated at the closing rate, with the difference in the "
     "translation reserve. Associate One reports in AED, which is pegged to the USD, so no exchange "
     "difference arises; for a floating currency it would also go to the translation reserve.",
     "IAS 21.39-41, 21.44, 21.47; IAS 28.10"),
    ("Combine line by line (subsidiaries only)",
     "The assets, liabilities, income and expenses of the parent and its subsidiaries are added together in full, "
     "even where the group owns less than 100%.", "IFRS 10.B86(a)"),
    ("Eliminate the investment; goodwill",
     "The parent's investment in each subsidiary is offset against its share of the subsidiary's equity at "
     "acquisition. Any excess is goodwill, which is not amortised but tested for impairment at least every year.",
     "IFRS 10.B86(b); IFRS 3.32; IAS 36.90"),
    ("Non-controlling interests",
     "NCI is measured at its proportionate share of the subsidiary's identifiable net assets (the IFRS 3 choice "
     "made here, so no goodwill is attributed to NCI). It is shown in equity, separately from the parent's "
     "equity, and receives its share of profit and of translation differences.", "IFRS 3.19; IFRS 10.22, B94"),
    ("Indirect holdings (sub-subsidiaries)",
     "Control passes down the chain, so a subsidiary of a subsidiary is the parent's subsidiary and is "
     "consolidated in full. Goodwill and NCI use the group's effective interest: Subsidiary Three 100% x 70% = 70%, "
     "Subsidiary Four 80% x 75% = 60%. NCI in Subsidiary Four is 40%: the 25% direct minority plus the "
     "Subsidiary Two minority's 20% x 75% indirect share, less that minority's 20% share of Subsidiary Two's cost "
     "of investment.", "IFRS 10.B86, Appendix A; IFRS 3.19, 3.32"),
    ("Foreign holding company",
     "Subsidiary Two holds Subsidiary Four in euros. Its investment is in its translated trial balance at the "
     "closing rate, so that amount is eliminated. Goodwill is measured with the cost at the historical rate, and "
     "the translation difference on the investment is reversed: the group's share from the translation reserve "
     "and the Subsidiary Two minority's share from NCI.", "IAS 21.39, 21.47; IFRS 10.B86(b)"),
    ("Intragroup transactions",
     "Intragroup balances, sales, purchases and unrealised profit in closing inventory are eliminated in full. "
     "The unrealised profit is charged to the selling entity; the seller is 100% owned, so NCI is not affected.",
     "IFRS 10.B86(c)"),
    ("Equity method (associates)",
     "An associate is not combined line by line. The investment is recognised at cost and then increased or "
     "decreased by the group's share of the associate's profit or loss after acquisition. Dividends received "
     "reduce the carrying amount instead of being income, so the parent's dividend income is reversed.",
     "IAS 28.10, 28.16"),
    ("Goodwill on associates",
     "The excess of cost over the group's share of the associate's net assets is included in the carrying "
     "amount of the investment. It is not shown separately, not amortised and not tested for impairment on "
     "its own.", "IAS 28.32, 28.42"),
    ("Transactions with associates",
     "Only the group's share of unrealised profit on sales between the group and an associate is eliminated. "
     "For a sale to the associate (downstream) it reduces the investment; for a purchase from the associate "
     "(upstream) it reduces group inventory. Balances with associates are not eliminated, because associates "
     "are outside the group.", "IAS 28.28-29"),
    ("Impairment of associates",
     "After applying the equity method, the group checks for objective evidence of impairment. If there is, the "
     "whole carrying amount is tested as a single asset under IAS 36 and any loss reduces the investment. "
     "Associate Two carries an impairment loss.", "IAS 28.40-43; IAS 36"),
    ("Losses of associates",
     "If the group's share of losses equals or exceeds its interest, it stops recognising further losses unless "
     "it has an obligation to fund them. Not triggered here.", "IAS 28.38-39"),
    ("Presentation",
     "Investments in associates and the share of their profit are separate line items. Profit and equity are "
     "split between owners of the parent and NCI.", "IAS 1.54(e), 1.82(c), 1.81B; IFRS 10.22"),
    ("Disclosure",
     "Summarised financial information for each material associate and a reconciliation to its carrying amount "
     "(shown on the Associates sheet).", "IFRS 12.21, B12-B14"),
    ("IFRS 18 (from 2027)",
     "IFRS 18 replaces IAS 1 for periods beginning on or after 1 January 2027. The share of profit of associates "
     "will then sit in the investing category, outside operating profit, as it already does in this "
     "income statement. These 2025 statements follow IAS 1.", "IFRS 18"),
]


def _ifrs_sheet(wb, entities, assoc, ent_row, acol, ar):
    ws = _sheet(wb, "IFRS Treatment", "How each entity is treated under IFRS",
                "Parent incorporated in the UAE; consolidated financial statements prepared under IFRS Accounting Standards")
    names = entities.set_index("entity")["name"]
    wrap = Alignment(wrap_text=True, vertical="top")
    _header(ws, 5, ["Entity", "Relationship", "Ownership", "Standards", "Method", "How it is treated in this workbook"],
            widths=[26, 34, 11, 22, 24, 80])
    for col in range(2, 8):
        ws.cell(row=5, column=col).alignment = Alignment(horizontal="right" if col == 4 else "left")
    rows = []
    for i, e in enumerate(entities.itertuples(index=False)):
        if i == 0:
            rows.append((e.name, "Parent and reporting entity, incorporated in the UAE", None, "IFRS 10, IAS 1, IAS 21",
                         "Prepares the consolidated statements",
                         "The UAE requires companies to prepare their financial statements under IFRS, so the group "
                         "reports under IFRS. Combined in full. Functional and presentation currency "
                         f"{GROUP_CCY}: most sales, pricing and financing are in {GROUP_CCY} (IAS 21.9), and the "
                         "UAE dirham is pegged to it."))
            continue
        why = "100% of the voting rights gives control" if e.ownership == 1 else \
            f"{e.ownership:.0%} of the voting rights (a majority) gives control"
        if e.held_by != entities.entity.iloc[0]:
            why = (f"{e.ownership:.0%} held by {names[e.held_by]}, so the parent controls it indirectly "
                   f"(sub-subsidiary)")
        how = ("Combined line by line. The investment is eliminated against share capital and pre-acquisition "
               "reserves; any excess of cost is goodwill.")
        std = "IFRS 10, IFRS 3"
        if e.currency != GROUP_CCY:
            std += ", IAS 21"; how = f"Translated from {e.currency} to {GROUP_CCY} first. " + how
        if e.held_by != entities.entity.iloc[0]:
            how += (f" Group effective interest is {e.effective:.0%}; goodwill and NCI are measured on that "
                    f"interest.")
            holder_ccy = entities.set_index("entity").currency[e.held_by]
            if holder_ccy != GROUP_CCY:
                how += (f" The {holder_ccy} holding company's investment is removed at the closing rate and its "
                        f"translation difference reversed.")
        if e.currency != GROUP_CCY:
            how += f" Goodwill is held in {e.currency} and retranslated at the closing rate."
        if e.effective < 1:
            how += f" The other {1 - e.effective:.0%} is a non-controlling interest, shown in equity, with its share of profit and translation differences."
        rows.append((e.name, f"Subsidiary: {why}", f"=Inputs!$E${ent_row[e.entity]}", std,
                     "Full consolidation (acquisition method)" + (" with NCI" if e.effective < 1 else ""), how))
    for a in assoc.itertuples(index=False) if len(assoc) else []:
        why = (f"{a.ownership:.0%} of the voting rights: presumed significant influence, not control"
               if a.ownership >= 0.2 else f"{a.ownership:.0%}: significant influence shown by board representation")
        how = ("Not combined line by line. One line on the balance sheet (investment at cost plus share of "
               "post-acquisition profits, less dividends received) and one line in the income statement (share of "
               "profit). Goodwill stays inside the investment. Only the group's share of unrealised profit on "
               "trading with the associate is eliminated.")
        std = "IAS 28" + (", IAS 21" if a.currency != GROUP_CCY else "")
        if a.impairment:
            how += " The whole investment was tested for impairment under IAS 36 and written down."; std += ", IAS 36"
        if a.currency != GROUP_CCY:
            how += f" Its results are translated from {a.currency}."
        rows.append((a.name, f"Associate: {why}", f"=Associates!{acol[a.entity]}{ar['own']}", std, "Equity method", how))
    for i, row in enumerate(rows):
        r = 6 + i
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=2 + j, value=v); c.font = _f(j == 0, color="008000" if j == 2 and v else "000000")
            c.alignment = wrap
            if j == 2: c.number_format = "0%"; c.alignment = Alignment(horizontal="right", vertical="top")
        ws.row_dimensions[r].height = 13 * max(2, len(row[5]) // 85 + 1) + 4
    r = 6 + len(rows) + 2
    ws.cell(row=r, column=2, value="Consolidation procedures and IFRS references").font = _f(True, size=12)
    _header(ws, r + 1, ["Step", "IFRS reference", "What the workbook does", "", "", ""])
    ws.merge_cells(start_row=r + 1, start_column=4, end_row=r + 1, end_column=7)
    for col in range(2, 8):
        ws.cell(row=r + 1, column=col).alignment = Alignment(horizontal="left")
    for i, (step, text, ref) in enumerate(IFRS_STEPS):
        rr = r + 2 + i
        ws.cell(row=rr, column=2, value=f"{i + 1}. {step}").font = _f(True)
        ws.cell(row=rr, column=3, value=ref).font = _f(color="1F4E78")
        ws.cell(row=rr, column=4, value=text).font = _f()
        ws.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=7)
        for col in (2, 3, 4): ws.cell(row=rr, column=col).alignment = wrap
        ws.row_dimensions[rr].height = 13 * max(2, len(text) // 125 + 1) + 4
    ws.cell(row=r + 3 + len(IFRS_STEPS), column=2, value="Paragraph references are to the IFRS Accounting Standards "
            "as issued by the IASB.").font = _f(italic=True, color="595959")


# ---------------------------------------------------------------- main
def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("--data", default="data")
    p.add_argument("--out", default="output/group_consolidation.xlsx")
    a = p.parse_args(argv)
    try:
        wide, coa, entities, fx, ic, assoc = load_inputs(a.data)
        usd, journal, elim, consolidated, em = consolidate(wide, coa, entities, fx, ic, assoc)
        s = summarise(usd, entities, consolidated)
        if abs(s["Total assets"] - s["Total equity and liabilities"]) > 0.01:
            raise ValidationError("Consolidated balance sheet does not balance")
    except ValidationError as e:
        print(f"Stopped: {e}"); return 1
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    write_workbook(wide, coa, entities, fx, ic, journal, a.out, assoc)
    print(f"Checks passed: {len(entities)} trial balances balance, {len(assoc)} associates equity-accounted, {len({j[0] for j in journal})} elimination entries "
          f"balance, intercompany fully eliminated, balance sheet balances.")
    for k, v in s.items():
        print(f"  {k:<32}{v:>10,.1f}")
    print(f"Saved {a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
